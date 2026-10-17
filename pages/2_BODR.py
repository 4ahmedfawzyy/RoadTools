import datetime
import io
import os
import uuid

import pandas as pd
import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

st.set_page_config(page_title="CO-BODR Engineering", layout="wide", page_icon="🛣️")

st.markdown(
    """
    <style>
        [data-testid="stSidebarNav"] { display: none !important; }
        .block-container { padding-top: 1.6rem; max-width: 1300px; }
        .hero { background: linear-gradient(135deg, #1F3864 0%, #2E5C9A 100%); padding: 1.5rem 2rem;
                border-radius: 14px; margin-bottom: 1.2rem; box-shadow: 0 4px 14px rgba(0,0,0,.18); }
        .hero-title { color: #fff; font-size: 1.9rem; font-weight: 700; line-height: 1.2; }
        .hero-sub { color: rgba(255,255,255,.85); font-size: .98rem; margin-top: .35rem; }
        .metric-card { background: rgba(128,128,128,.08); border: 1px solid rgba(128,128,128,.25);
                       border-radius: 12px; padding: 1rem 1.2rem; text-align: center; }
        .metric-label { font-size: .8rem; font-weight: 600; letter-spacing: .06em; text-transform: uppercase; opacity: .7; }
        .metric-value { font-size: 2rem; font-weight: 800; color: #5B9BE8; margin-top: .1rem; }
        .note { font-size: .88rem; opacity: .75; }
        button[data-baseweb="tab"] { font-weight: 600; }
        .disclaimer { font-size: .8rem; opacity: .65; text-align: center; margin: 2rem 0 .5rem; }
    </style>
    """,
    unsafe_allow_html=True,
)

# ----------------------------------------------------------------------------
# Data (AASHTO)
# ----------------------------------------------------------------------------
turning_data = {
    "Design Vehicle Type": ["Passenger Car", "Single-Unit Truck", "Intercity Bus", "WB-12 Semitrailer", "WB-20 Semitrailer"],
    "Symbol": ["P", "SU-12", "BUS-12", "WB-12", "WB-20"],
    "Min Design Turning Radius (m)": [7.26, 15.60, 12.70, 12.16, 13.66],
    "Centerline Radius (CTR) (m)": [6.40, 14.46, 11.53, 10.97, 12.50],
    "Min Inside Radius (m)": [4.39, 11.09, 7.41, 5.88, 0.59],
}
df_turning = pd.DataFrame(turning_data)
vehicle_images = {
    "P": "Passenger Car Turning Radius.png", "SU-12": "SU-12 Turning Radius.png",
    "BUS-12": "BUS-12 Turning Radius.png", "WB-12": "WB-12 Turning Radius.png",
    "WB-20": "WB-20 Turning Radius.png",
}

k_values_data = {
    "Crest": {20: 1, 30: 2, 40: 4, 50: 7, 60: 11, 70: 17, 80: 26, 90: 39, 100: 52, 110: 74, 120: 95, 130: 124},
    "Sag": {20: 3, 30: 6, 40: 9, 50: 13, 60: 18, 70: 23, 80: 30, 90: 38, 100: 45, 110: 55, 120: 63, 130: 73},
}
df_k = pd.DataFrame({
    "Design Speed (km/h)": list(k_values_data["Crest"].keys()),
    "Crest K": list(k_values_data["Crest"].values()),
    "Sag K": list(k_values_data["Sag"].values()),
})

horizontal_data = {
    "e (%)": [-6.0, -5.0, -4.0, -3.0, -2.0, -1.5, 0.0, 1.5, 2.0, 3.0, 4.0, 5.0, 6.0, 7.0, 8.0, 9.0, 10.0, 11.0, 12.0],
    "20": [11, 10, 10, 10, 10, 9, 9, 9, 9, 8, 8, 8, 8, 7, 7, 7, 7, 7, 7],
    "30": [32, 31, 30, 28, 27, 27, 25, 24, 24, 23, 22, 21, 21, 20, 20, 19, 19, 18, 18],
    "40": [74, 70, 66, 63, 60, 59, 55, 51, 50, 48, 47, 45, 43, 42, 41, 39, 38, 37, 36],
    "50": [151, 141, 131, 123, 116, 113, 104, 96, 94, 89, 86, 82, 79, 76, 73, 70, 68, 66, 64],
    "60": [258, 236, 218, 202, 189, 183, 167, 153, 149, 142, 135, 129, 123, 118, 113, 109, 105, 101, 98],
    "70": [429, 386, 351, 322, 297, 286, 257, 234, 227, 214, 203, 193, 184, 175, 168, 161, 154, 148, 143],
}
df_horizontal = pd.DataFrame(horizontal_data)
H_SPEEDS = [c for c in df_horizontal.columns if c != "e (%)"]

STRONG = "background-color: #dc3545; color: #ffffff; font-weight: 700"
LIGHT = "background-color: rgba(220,53,69,.18)"


def highlight(df, row_mask, strong_cols=()):
    """Light red on the matching row, solid red on the chosen cell(s)."""
    def style(row):
        if not row_mask(row):
            return [""] * len(row)
        return [STRONG if c in strong_cols else LIGHT for c in row.index]
    return df.style.apply(style, axis=1)


# ----------------------------------------------------------------------------
# Report model
# ----------------------------------------------------------------------------
if "report_data" not in st.session_state:
    st.session_state.report_data = []


def add_item(item):
    for existing in st.session_state.report_data:
        if {k: v for k, v in existing.items() if k != "id"} == item:
            st.toast("Already in the report.", icon="ℹ️")
            return
    st.session_state.report_data.append({"id": uuid.uuid4().hex, **item})
    st.toast("Added to the calculation note.", icon="✅")


def remove_item(item_id):
    st.session_state.report_data = [i for i in st.session_state.report_data if i["id"] != item_id]


def clear_items():
    st.session_state.report_data = []


def describe(item):
    if item["tool"] == "K-Value":
        return (f"📈 **K-value** · {item['curve_type']} · {item['speed']} km/h",
                f"Required minimum K = **{item['k_value']}**")
    if item["tool"] == "Horizontal Radius":
        return (f"🔄 **Horizontal radius** · Vd {item['speed']} km/h · e = {item['e_val']}%",
                f"Minimum radius R = **{item['radius']} m**")
    return (f"📐 **Turning radius** · {item['vehicle_name']} ({item['symbol']})",
            f"Min design radius **{item['min_r']} m** · CTR {item['ctr']} m · inside radius {item['inside_r']} m")


# ----------------------------------------------------------------------------
# Word calculation note
# ----------------------------------------------------------------------------
def _shade(cell, fill):
    cell._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls("w"), fill)))


def _fill(cell, text, bold=False, size=9, color=None, center=True):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _table(doc, headers, rows, strong=(), light_rows=()):
    """strong: (row, col) cells in solid red; light_rows: row indexes in light red."""
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for c, h in enumerate(headers):
        _fill(t.rows[0].cells[c], h, bold=True, color="FFFFFF")
        _shade(t.rows[0].cells[c], "1F3864")
    for r, row in enumerate(rows):
        cells = t.add_row().cells
        for c, val in enumerate(row):
            is_strong = (r, c) in strong
            _fill(cells[c], val, bold=is_strong, color="FFFFFF" if is_strong else None)
            if is_strong:
                _shade(cells[c], "DC3545")
            elif r in light_rows:
                _shade(cells[c], "F8D7DA")
    return t


def build_docx(items, project="", prepared_by=""):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading("CO-BODR - Calculation Note", 0)

    info = doc.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    for label, value in (("Project", project or "-"), ("Prepared by", prepared_by or "-"),
                         ("Date", datetime.date.today().strftime("%d %B %Y")), ("Design code", "AASHTO")):
        cells = info.add_row().cells
        _fill(cells[0], label, bold=True, size=10, center=False)
        _shade(cells[0], "E7EEF8")
        _fill(cells[1], value, size=10, center=False)
    doc.add_paragraph()

    if not items:
        doc.add_paragraph("No data added to the report yet.")
        return doc

    n = 0
    ks = [i for i in items if i["tool"] == "K-Value"]
    hs = [i for i in items if i["tool"] == "Horizontal Radius"]
    ts = [i for i in items if i["tool"] == "Turning Radius"]

    if ks:
        n += 1
        doc.add_heading(f"{n}. Vertical Alignment (K-Value)", level=1)
        _table(doc, ["Curve type", "Design speed (km/h)", "Required minimum K"],
               [[i["curve_type"], i["speed"], i["k_value"]] for i in ks])
        doc.add_paragraph()
        doc.add_heading("Reference table (selected values highlighted)", level=2)
        speeds = list(df_k["Design Speed (km/h)"])
        rows = [[r["Design Speed (km/h)"], r["Crest K"], r["Sag K"]] for _, r in df_k.iterrows()]
        strong = {(speeds.index(i["speed"]), 1 if i["curve_type"] == "Crest" else 2) for i in ks}
        _table(doc, ["Design speed (km/h)", "Crest K", "Sag K"], rows, strong=strong,
               light_rows={speeds.index(i["speed"]) for i in ks})
        doc.add_paragraph()

    if hs:
        n += 1
        doc.add_heading(f"{n}. Horizontal Alignment (Low-Speed Urban)", level=1)
        _table(doc, ["Design speed (km/h)", "Superelevation e (%)", "Minimum radius R (m)"],
               [[i["speed"], i["e_val"], i["radius"]] for i in hs])
        doc.add_paragraph()
        doc.add_heading("AASHTO Table 3-13a (selected values highlighted)", level=2)
        es = [float(e) for e in df_horizontal["e (%)"]]
        cols = list(df_horizontal.columns)
        rows = [[r[c] for c in cols] for _, r in df_horizontal.iterrows()]
        strong = {(es.index(float(i["e_val"])), cols.index(str(i["speed"]))) for i in hs}
        _table(doc, ["e (%)"] + [f"Vd = {c} km/h" for c in cols[1:]], rows, strong=strong,
               light_rows={es.index(float(i["e_val"])) for i in hs})
        doc.add_paragraph()

    if ts:
        n += 1
        doc.add_heading(f"{n}. Turning Radius Check", level=1)
        _table(doc, ["Design vehicle", "Symbol", "Min design radius (m)", "CTR (m)", "Min inside radius (m)"],
               [[i["vehicle_name"], i["symbol"], i["min_r"], i["ctr"], i["inside_r"]] for i in ts])
        doc.add_paragraph()
        for sym in dict.fromkeys(i["symbol"] for i in ts):
            path = vehicle_images.get(sym)
            if path and os.path.exists(path):
                doc.add_paragraph(f"Design vehicle {sym}:")
                doc.add_picture(path, width=Inches(5.5))

    note = doc.add_paragraph("Preliminary values taken from the tool tables. Verify against the governing design code.")
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(9)
    return doc


def docx_bytes(items, project, prepared_by):
    bio = io.BytesIO()
    build_docx(items, project, prepared_by).save(bio)
    return bio.getvalue()


# ----------------------------------------------------------------------------
# Sidebar
# ----------------------------------------------------------------------------
if st.sidebar.button("🏠 Home", use_container_width=True, type="primary"):
    st.switch_page("app.py")
st.sidebar.markdown("## 🛣️ CO-BODR")
st.sidebar.metric("Items in report", len(st.session_state.report_data))
st.sidebar.caption("Add results with the ➕ buttons, then open the **Report** tab to download the calculation note.")

# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
st.markdown(
    '<div class="hero"><div class="hero-title">🛣️ CO-BODR · Highway Design Wizard</div>'
    '<div class="hero-sub">AASHTO design values for vertical curves, horizontal curves and turning radii, '
    'with a downloadable calculation note</div></div>',
    unsafe_allow_html=True,
)


def metric(label, value):
    return f'<div class="metric-card"><div class="metric-label">{label}</div><div class="metric-value">{value}</div></div>'


tab_k, tab_h, tab_t, tab_r = st.tabs([
    "📈 Vertical Curves (K)", "🔄 Horizontal Curves (R)", "📐 Turning Radius",
    f"🧾 Report ({len(st.session_state.report_data)})",
])

# ---- Vertical curves -----------------------------------------------------------
with tab_k:
    c1, c2 = st.columns(2)
    curve_type = c1.selectbox("Curve type", ["Crest", "Sag"])
    speed = c2.selectbox("Design speed [km/h]", list(k_values_data[curve_type].keys()), index=10)
    req_k = k_values_data[curve_type][speed]

    st.markdown(metric("Required minimum K-value", req_k), unsafe_allow_html=True)
    st.write("")
    if st.button("➕ Add to report", type="primary", key="add_k"):
        add_item({"tool": "K-Value", "curve_type": curve_type, "speed": int(speed), "k_value": int(req_k)})

    st.markdown("##### Reference table")
    st.dataframe(highlight(df_k, lambda r: r["Design Speed (km/h)"] == speed, [f"{curve_type} K"]),
                 hide_index=True, use_container_width=True, height=35 * (len(df_k) + 1) + 3)

# ---- Horizontal curves ---------------------------------------------------------
with tab_h:
    st.caption("Based on AASHTO Table 3-13a: minimum radii for low-speed urban streets (up to 70 km/h).")
    c1, c2 = st.columns(2)
    v_design = c1.selectbox("Design speed (Vd) [km/h]", H_SPEEDS, index=3)
    e_val = c2.selectbox("Superelevation (e) [%]", df_horizontal["e (%)"].tolist(), index=6)
    radius = int(df_horizontal.loc[df_horizontal["e (%)"] == e_val, v_design].values[0])

    st.markdown(metric("Minimum curve radius (R)", f"{radius} m"), unsafe_allow_html=True)
    st.write("")
    if st.button("➕ Add to report", type="primary", key="add_h"):
        add_item({"tool": "Horizontal Radius", "speed": v_design, "e_val": str(e_val), "radius": radius})

    st.markdown("##### AASHTO Table 3-13a")
    show = df_horizontal.rename(columns={c: f"{c} km/h" for c in H_SPEEDS})
    st.dataframe(
        highlight(show, lambda r: r["e (%)"] == e_val, [f"{v_design} km/h"]).format({"e (%)": "{:.1f}"}),
        hide_index=True, use_container_width=True, height=35 * (len(show) + 1) + 3)

# ---- Turning radius ------------------------------------------------------------
with tab_t:
    names = df_turning["Design Vehicle Type"] + " (" + df_turning["Symbol"] + ")"
    selected_display = st.selectbox("Design vehicle", names)
    selected_symbol = selected_display.split("(")[-1].replace(")", "")
    v = df_turning[df_turning["Symbol"] == selected_symbol].iloc[0]

    m1, m2, m3 = st.columns(3)
    m1.markdown(metric("Min design radius", f"{v['Min Design Turning Radius (m)']} m"), unsafe_allow_html=True)
    m2.markdown(metric("Centerline radius (CTR)", f"{v['Centerline Radius (CTR) (m)']} m"), unsafe_allow_html=True)
    m3.markdown(metric("Min inside radius", f"{v['Min Inside Radius (m)']} m"), unsafe_allow_html=True)
    st.write("")
    if st.button("➕ Add to report", type="primary", key="add_t"):
        add_item({"tool": "Turning Radius", "vehicle_name": v["Design Vehicle Type"], "symbol": selected_symbol,
                  "min_r": float(v["Min Design Turning Radius (m)"]), "ctr": float(v["Centerline Radius (CTR) (m)"]),
                  "inside_r": float(v["Min Inside Radius (m)"])})

    img = vehicle_images.get(selected_symbol)
    if img and os.path.exists(img):
        st.image(img, caption=f"{selected_symbol} turning path", width=620)
    st.markdown("##### Design vehicles")
    st.dataframe(highlight(df_turning, lambda r: r["Symbol"] == selected_symbol,
                           ["Min Design Turning Radius (m)"]).format(precision=2),
                 hide_index=True, use_container_width=True)

# ---- Report ----------------------------------------------------------------------
with tab_r:
    items = st.session_state.report_data
    p1, p2 = st.columns(2)
    project = p1.text_input("Project name", key="proj_name")
    prepared_by = p2.text_input("Prepared by", key="proj_by")

    if not items:
        st.info("Nothing in the report yet. Use the ➕ Add to report buttons in the other tabs.")
    else:
        for it in items:
            title, detail = describe(it)
            with st.container(border=True):
                a, b = st.columns([8, 1])
                a.markdown(f"{title}  \n{detail}")
                b.button("🗑️", key=f"rm_{it['id']}", on_click=remove_item, args=(it["id"],),
                         help="Remove from report", use_container_width=True)
        d1, d2 = st.columns([3, 1])
        try:
            d1.download_button("⬇️ Download calculation note (Word)", docx_bytes(items, project, prepared_by),
                               "Calculation_Note.docx",
                               "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               type="primary", use_container_width=True)
        except Exception as e:  # keep the page usable if the note cannot be built
            d1.error(f"Could not build the Word file: {e}")
        d2.button("Clear all", on_click=clear_items, use_container_width=True)

st.markdown('<div class="disclaimer">⚠️ Preliminary values from the tool tables. '
            'Verify against the governing design code.</div>', unsafe_allow_html=True)
