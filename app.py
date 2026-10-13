import streamlit as st
import pandas as pd
import io
import os
import datetime
from docx import Document
from docx.shared import Inches

# ----------------- إعدادات الصفحة -----------------
st.set_page_config(page_title="CO-BODR Engineering", layout="wide", page_icon="logo.png")

# ----------------- تنسيقات CSS -----------------
st.markdown("""
    <style>
    .result-box { background-color: #f1f3f5; padding: 15px; border-radius: 8px; text-align: center; margin-top: 10px; border: 1px solid #dee2e6; }
    .result-value { font-size: 24px; font-weight: bold; color: #007bff; display: block; margin-top: 5px; }
    .cs-card { background-color: #f8f9fa; border: 1px solid #dee2e6; padding: 12px; border-radius: 5px; margin-bottom: 10px; color: #333; font-weight: 500; }
    </style>
""", unsafe_allow_html=True)

# ----------------- تهيئة الذاكرة -----------------
if 'cross_sections' not in st.session_state:
    st.session_state.cross_sections = []

# ----------------- الداتا الخاصة بالأكواد -----------------
data = {
    "AASHTO": {
        "k_crest": {20: 1, 30: 2, 40: 4, 50: 7, 60: 11, 70: 17, 80: 26, 90: 39, 100: 52, 110: 74, 120: 95, 130: 124},
        "k_sag": {20: 3, 30: 6, 40: 9, 50: 13, 60: 18, 70: 23, 80: 30, 90: 38, 100: 45, 110: 55, 120: 63, 130: 73},
        "vehicles": pd.DataFrame({'Name': ['Passenger Car', 'Single-Unit Truck', 'Intercity Bus', 'WB-12', 'WB-20'], 'Symbol': ['P', 'SU-12', 'BUS-12', 'WB-12', 'WB-20'], 'Min_R': [7.26, 15.60, 12.70, 12.16, 13.66], 'Inside_R': [4.39, 11.09, 7.41, 5.88, 0.59]})
    },
    "Abu Dhabi (ADM)": {
        "k_crest": {20: 1, 30: 2, 40: 4, 50: 7, 60: 11, 70: 17, 80: 26, 90: 39, 100: 52, 110: 74, 120: 95, 130: 124, 140: 161},
        "k_sag": {20: 3, 30: 6, 40: 9, 50: 13, 60: 18, 70: 23, 80: 30, 90: 38, 100: 45, 110: 55, 120: 63, 130: 73, 140: 84},
        "vehicles": pd.DataFrame({'Name': ['Passenger Car', 'Single Unit SU-9', 'Single Unit SU-12', 'City Bus', 'WB-12', 'WB-15', 'WB-20'], 'Symbol': ['P', 'SU-9', 'SU-12', 'CITY-BUS', 'WB-12', 'WB-15', 'WB-20'], 'Min_R': [7.26, 12.73, 15.60, 12.50, 12.16, 13.66, 13.66], 'Inside_R': [4.39, 8.64, 11.09, 5.30, 5.88, 5.20, 0.59]})
    },
    "MOMRA": {
        "k_crest": {30: 2, 40: 4, 50: 7, 60: 11, 70: 17, 80: 26, 90: 39, 100: 52, 110: 74, 120: 95, 130: 124},
        "k_sag": {30: 6, 40: 9, 50: 13, 60: 18, 70: 23, 80: 30, 90: 38, 100: 45, 110: 55, 120: 63, 130: 73},
        "vehicles": pd.DataFrame({'Name': ['سيارة ركوب صغيرة', 'شاحنة وحدة مفردة', 'حافلة', 'شاحنة بمقطورة متوسطة', 'شاحنة بمقطورة كبيرة'], 'Symbol': ['P', 'SU', 'BUS-12', 'WB-12', 'WB-20'], 'Min_R': [7.26, 12.73, 12.70, 12.16, 13.66], 'Inside_R': [4.39, 8.64, 7.41, 5.88, 0.59]})
    }
}

# خرائط الصور المرجعية
vehicle_images_map = {
    'P': ['Min. Turning Radius Passenger Car.png', 'Passenger Car Turning Radius.png'],
    'SU': ['Min. Turning Radius SU.png'], 'SU-9': ['Min. Turning Radius SU.png'], 'SU-12': ['Min. Turning Radius SU.png', 'SU-12 Turning Radius.png'],
    'CITY-BUS': ['Min. Turning Radius BUS-12.png'], 'BUS-12': ['Min. Turning Radius BUS-12.png', 'BUS-12 Turning Radius.png'],
    'WB-12': ['Min. Turning Radius WB-12.png', 'WB-12 Turning Radius.png'],
    'WB-15': ['Min. Turning Radius WB-20.png'], 'WB-20': ['Min. Turning Radius WB-20.png', 'WB-20 Turning Radius.png'], 'WB-33D': ['Min. Turning Radius WB-20.png']
}

k_images_map = {
    "Abu Dhabi (ADM)": {"crest": "K-Value Crest.png", "sag": "K-Value Sag.png"},
    "MOMRA": {"crest": "Min. K-Value Crest.png", "sag": "Min. K-Value Sag.png"}
}

st.title("🛣️ CO-BODR Highway Design Wizard")

# ================= STEP 1: SELECT CODE =================
with st.container(border=True):
    st.markdown("### 1. Select Design Code (Reference Standard)")
    selected_code = st.radio("Standard:", ["AASHTO", "Abu Dhabi (ADM)", "MOMRA"], horizontal=True, label_visibility="collapsed")
    active_data = data[selected_code]

# ================= STEP 2: SELECT SPEED =================
with st.container(border=True):
    st.markdown("### 2. Vertical Alignment (K-Values)")
    speed_options = list(active_data["k_crest"].keys())
    selected_speed = st.selectbox("Design Speed (km/h):", speed_options)
    
    req_k_crest = active_data["k_crest"].get(selected_speed, "N/A")
    req_k_sag = active_data["k_sag"].get(selected_speed, "N/A")
    
    c1, c2 = st.columns(2)
    c1.markdown(f'<div class="result-box">Required Crest K-Value<span class="result-value">{req_k_crest}</span></div>', unsafe_allow_html=True)
    c2.markdown(f'<div class="result-box">Required Sag K-Value<span class="result-value">{req_k_sag}</span></div>', unsafe_allow_html=True)

# ================= STEP 3: SELECT VEHICLE =================
with st.container(border=True):
    st.markdown("### 3. Vehicle Turning Radius")
    df_v = active_data["vehicles"]
    display_names = df_v['Name'] + " (" + df_v['Symbol'] + ")"
    selected_v = st.selectbox("Design Vehicle Type:", display_names)
    
    v_symbol = selected_v.split("(")[-1].replace(")", "")
    v_row = df_v[df_v['Symbol'] == v_symbol].iloc[0]
    
    c1, c2 = st.columns(2)
    c1.markdown(f'<div class="result-box">Minimum Design Radius<span class="result-value">{v_row["Min_R"]} m</span></div>', unsafe_allow_html=True)
    c2.markdown(f'<div class="result-box">Minimum Inside Radius<span class="result-value">{v_row["Inside_R"]} m</span></div>', unsafe_allow_html=True)

# ================= STEP 4: TYPICAL CROSS SECTION =================
with st.container(border=True):
    st.markdown("### 4. Typical Cross Sections Builder")
    cs_name = st.text_input("📝 Cross Section Name (e.g., Main Road):", value="Typical Section 1")
    
    col_left, col_mid, col_right = st.columns([1, 0.4, 1])
    with col_left:
        st.markdown("#### ⬅️ Left Direction")
        l_lanes = st.number_input("Drive Lanes (Count)", min_value=0, max_value=6, value=2, key="ll")
        l_width = st.number_input("Lane Width (m)", min_value=2.0, max_value=4.0, value=3.3, step=0.1, key="lw")
        l_park = st.number_input("Parking Lane (m)", min_value=0.0, max_value=5.0, value=0.0, step=0.1, key="lp")
        l_sw = st.number_input("Sidewalk/Shoulder (m)", min_value=0.0, max_value=10.0, value=2.0, step=0.5, key="ls")
        
    with col_mid:
        st.markdown("#### 🟩 Median")
        median = st.number_input("Median Width (m)", min_value=0.0, max_value=20.0, value=2.0, step=0.5, key="med")
        
    with col_right:
        st.markdown("#### ➡️ Right Direction")
        r_lanes = st.number_input("Drive Lanes (Count) ", min_value=0, max_value=6, value=2, key="rl")
        r_width = st.number_input("Lane Width (m) ", min_value=2.0, max_value=4.0, value=3.3, step=0.1, key="rw")
        r_park = st.number_input("Parking Lane (m) ", min_value=0.0, max_value=5.0, value=0.0, step=0.1, key="rp")
        r_sw = st.number_input("Sidewalk/Shoulder (m) ", min_value=0.0, max_value=10.0, value=2.0, step=0.5, key="rs")

    l_total = (l_lanes * l_width)
    r_total = (r_lanes * r_width)
    total_row = l_sw + l_park + l_total + median + r_total + r_park + r_sw

    def build_html_block(width, color, label):
        if width <= 0: return ""
        return f'<div style="flex: {width}; background-color: {color}; display: flex; align-items: center; justify-content: center; border: 1px solid rgba(255,255,255,0.2);">{label}<br>{round(width,2)}m</div>'

    # دمج الـ HTML بدون مسافات عشان ميتحولش لـ Code Block
    diagram_html = '<div style="display: flex; width: 100%; height: 70px; text-align: center; color: white; font-weight: bold; font-family: sans-serif; border-radius: 8px; overflow: hidden; margin-top: 15px; margin-bottom: 15px; font-size: 14px;">' + build_html_block(l_sw, '#95a5a6', 'SW') + build_html_block(l_park, '#f39c12', 'Park') + build_html_block(l_total, '#34495e', 'Lanes') + build_html_block(median, '#2ecc71', 'Med') + build_html_block(r_total, '#34495e', 'Lanes') + build_html_block(r_park, '#f39c12', 'Park') + build_html_block(r_sw, '#95a5a6', 'SW') + '</div>'
    
    st.markdown(diagram_html, unsafe_allow_html=True)
    st.markdown(f"**Total Right-of-Way (ROW) Width: <span style='color:#007bff; font-size: 22px;'>{round(total_row, 2)} m</span>**", unsafe_allow_html=True)

    if st.button("➕ Save Cross Section to Project"):
        st.session_state.cross_sections.append({
            'name': cs_name, 'l_lanes': l_lanes, 'l_width': l_width, 'l_park': l_park, 'l_sw': l_sw,
            'median': median, 'r_lanes': r_lanes, 'r_width': r_width, 'r_park': r_park, 'r_sw': r_sw, 'total_row': total_row
        })
        st.success(f"Section '{cs_name}' saved!")

    if st.session_state.cross_sections:
        st.markdown("#### 📑 Saved Sections for Report:")
        for idx, cs in enumerate(st.session_state.cross_sections):
            st.markdown(f"<div class='cs-card'><b>{idx+1}. {cs['name']}</b> | ROW = {round(cs['total_row'], 2)} m</div>", unsafe_allow_html=True)
        if st.button("🗑️ Clear All Sections"):
            st.session_state.cross_sections = []
            st.rerun()

# ================= STEP 5: GENERATE REPORT =================
def create_word_report(p_name, d_name):
    doc = Document()
    title = doc.add_heading('Highway Design Calculation Note', 0)
    title.alignment = 1 
    doc.add_paragraph(f"Project Name: {p_name}\nDesigned By: {d_name}\nDate: {datetime.datetime.now().strftime('%Y-%m-%d')}\nCode Reference: {selected_code}").alignment = 1
    doc.add_page_break()
    
    # 1. Vertical Alignment
    doc.add_heading('1. Vertical Alignment (K-Values)', level=1)
    table_k = doc.add_table(rows=2, cols=3)
    table_k.style = 'Table Grid'
    hdr_cells = table_k.rows[0].cells
    hdr_cells[0].text, hdr_cells[1].text, hdr_cells[2].text = 'Design Speed', 'Crest K-Value', 'Sag K-Value'
    row_cells = table_k.rows[1].cells
    row_cells[0].text, row_cells[1].text, row_cells[2].text = f"{selected_speed} km/h", str(req_k_crest), str(req_k_sag)
    
    if selected_code in k_images_map:
        doc.add_paragraph("\nReference Tables from Code:")
        crest_img = k_images_map[selected_code]["crest"]
        sag_img = k_images_map[selected_code]["sag"]
        if os.path.exists(crest_img): doc.add_picture(crest_img, width=Inches(5.0))
        if os.path.exists(sag_img): doc.add_picture(sag_img, width=Inches(5.0))
    doc.add_paragraph("_" * 50)
    
    # 2. Vehicle Turning Radius
    doc.add_heading('2. Vehicle Turning Radius', level=1)
    table_v = doc.add_table(rows=2, cols=3)
    table_v.style = 'Table Grid'
    hdr_cells = table_v.rows[0].cells
    hdr_cells[0].text, hdr_cells[1].text, hdr_cells[2].text = 'Design Vehicle', 'Min Design Radius (m)', 'Min Inside Radius (m)'
    row_cells = table_v.rows[1].cells
    row_cells[0].text, row_cells[1].text, row_cells[2].text = f"{v_row['Name']} ({v_row['Symbol']})", str(v_row['Min_R']), str(v_row['Inside_R'])
    
    img_names = vehicle_images_map.get(v_symbol, [])
    for img_name in img_names:
        if os.path.exists(img_name):
            doc.add_paragraph("\nVehicle Turning Path Reference:")
            doc.add_picture(img_name, width=Inches(5.5))
            break 
    doc.add_paragraph("_" * 50)
    
    # 3. Cross Sections
    doc.add_heading('3. Typical Cross Sections Details', level=1)
    if not st.session_state.cross_sections:
        doc.add_paragraph("No cross sections were added.")
    else:
        for cs in st.session_state.cross_sections:
            doc.add_heading(f"Section: {cs['name']}", level=2)
            table_cs = doc.add_table(rows=4, cols=2)
            table_cs.style = 'Table Grid'
            table_cs.rows[0].cells[0].text, table_cs.rows[0].cells[1].text = "Left Direction", f"{cs['l_lanes']} Lanes (@ {cs['l_width']}m), Parking: {cs['l_park']}m, SW: {cs['l_sw']}m"
            table_cs.rows[1].cells[0].text, table_cs.rows[1].cells[1].text = "Median", f"{cs['median']} m"
            table_cs.rows[2].cells[0].text, table_cs.rows[2].cells[1].text = "Right Direction", f"{cs['r_lanes']} Lanes (@ {cs['r_width']}m), Parking: {cs['r_park']}m, SW: {cs['r_sw']}m"
            table_cs.rows[3].cells[0].text, table_cs.rows[3].cells[1].text = "Total ROW Width", f"{round(cs['total_row'], 2)} m"
            doc.add_paragraph("")
            
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

with st.container(border=True):
    st.markdown("### 5. Generate Professional Calculation Note")
    col1, col2 = st.columns(2)
    proj_name = col1.text_input("Project Name:", value="Highway Design Project")
    designer_name = col2.text_input("Designed By:", value="Eng. Ahmed Fawzy Aamer Khalifa")
    
    st.download_button(
        label="📄 Download Word Report",
        data=create_word_report(proj_name, designer_name),
        file_name="CO_BODR_Report.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary",
        use_container_width=True
    )
